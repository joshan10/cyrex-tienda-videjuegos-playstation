from datetime import datetime

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.crud.ventas import (
    actualizar_venta,
    crear_venta,
    generar_datos_excel,
    obtener_stats_ventas,
    obtener_venta_por_id,
    obtener_venta_por_numero,
    obtener_ventas,
    reporte_ventas_detallado,
    ventas_por_fecha,
)
from app.dependencies import current_user, require_roles
from app.exceptions import RecursoNoEncontrado
from app.pagination import Paginacion, get_paginacion, paginate_query
from app.models.entities import Orden, Venta
from app.schemas.ventas import FacturaFiltros, VentaActualizar, VentaEntrada
from app.services.pdf_service import invoice_pdf, sales_report_pdf

router = APIRouter(prefix="/ventas", tags=["ventas"])


@router.post(
    "",
    summary="Crear venta (factura) desde una orden",
    status_code=201,
    responses={201: {"description": "Venta creada"}, 409: {"description": "Orden ya facturada o estado inválido"}},
)
def create_venta(
    data: VentaEntrada,
    _: dict = Depends(require_roles("Administrador", "Empleado")),
    db: Session = Depends(get_db),
):
    venta = crear_venta(
        db,
        orden_id=data.orden_id,
        impuesto_porcentaje=data.impuesto_porcentaje,
        descuento_porcentaje=data.descuento_porcentaje,
        metodo_pago=data.metodo_pago,
        notas_factura=data.notas_factura,
    )
    return {"message": "Venta registrada exitosamente", "venta": venta}


@router.get(
    "",
    summary="Listar ventas con filtros",
    responses={200: {"description": "Lista de ventas"}},
)
def list_ventas(
    numero_factura: str | None = Query(default=None),
    cliente_correo: str | None = Query(default=None),
    cliente_documento: str | None = Query(default=None),
    fecha_desde: str | None = Query(default=None, description="Formato: YYYY-MM-DD"),
    fecha_hasta: str | None = Query(default=None, description="Formato: YYYY-MM-DD"),
    estado: str | None = Query(default=None),
    producto_id: int | None = Query(default=None),
    servicio_id: int | None = Query(default=None),
    valor_minimo: float | None = Query(default=None),
    valor_maximo: float | None = Query(default=None),
    metodo_pago: str | None = Query(default=None),
    paginacion: Paginacion = Depends(get_paginacion),
    _: dict = Depends(require_roles("Administrador", "Empleado")),
    db: Session = Depends(get_db),
):
    filtros = FacturaFiltros(
        numero_factura=numero_factura,
        cliente_correo=cliente_correo,
        cliente_documento=cliente_documento,
        fecha_desde=fecha_desde,
        fecha_hasta=fecha_hasta,
        estado=estado,
        producto_id=producto_id,
        servicio_id=servicio_id,
        valor_minimo=valor_minimo,
        valor_maximo=valor_maximo,
        metodo_pago=metodo_pago,
    )
    items, total = obtener_ventas(db, filtros, paginacion.skip, paginacion.size)
    return {"items": items, **paginate_query(total, paginacion)}


@router.get(
    "/stats/dashboard",
    summary="Estadísticas de ventas para dashboard",
    responses={200: {"description": "Estadísticas"}},
)
def stats_dashboard(
    _: dict = Depends(require_roles("Administrador")),
    db: Session = Depends(get_db),
):
    return {"stats": obtener_stats_ventas(db)}


@router.get(
    "/mis-facturas",
    summary="Listar facturas del cliente autenticado",
)
def mis_facturas(user: dict = Depends(current_user), db: Session = Depends(get_db)):
    ventas = db.scalars(
        select(Venta)
        .join(Orden, Venta.orden_id == Orden.id)
        .where(Orden.usuario_id == user["id"])
        .order_by(Venta.fecha_venta.desc())
    ).all()
    return {"items": [obtener_venta_por_numero(db, venta.numero_factura) for venta in ventas]}


@router.get(
    "/reporte/diario",
    summary="Reporte diario de ventas",
    responses={200: {"description": "Reporte por día"}},
)
def reporte_diario(
    fecha_inicio: str | None = Query(default=None, description="Formato: YYYY-MM-DD"),
    fecha_fin: str | None = Query(default=None, description="Formato: YYYY-MM-DD"),
    _: dict = Depends(require_roles("Administrador")),
    db: Session = Depends(get_db),
):
    return {"reporte": ventas_por_fecha(db, fecha_inicio, fecha_fin)}


@router.get(
    "/reporte/detallado",
    summary="Reporte detallado de ventas con top productos",
    responses={200: {"description": "Reporte detallado"}},
)
def reporte_detallado(
    fecha_inicio: str | None = Query(default=None, description="Formato: YYYY-MM-DD"),
    fecha_fin: str | None = Query(default=None, description="Formato: YYYY-MM-DD"),
    _: dict = Depends(require_roles("Administrador")),
    db: Session = Depends(get_db),
):
    return {"reporte": reporte_ventas_detallado(db, fecha_inicio, fecha_fin)}


@router.get(
    "/reporte/excel",
    summary="Exportar reporte de ventas en Excel",
    responses={200: {"description": "Archivo Excel"}},
)
def reporte_excel(
    fecha_inicio: str | None = Query(default=None, description="Formato: YYYY-MM-DD"),
    fecha_fin: str | None = Query(default=None, description="Formato: YYYY-MM-DD"),
    _: dict = Depends(require_roles("Administrador")),
    db: Session = Depends(get_db),
):
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

    datos = generar_datos_excel(db, fecha_inicio, fecha_fin)

    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte de Ventas"

    # Estilos
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="171B24", end_color="171B24", fill_type="solid")
    accent_fill = PatternFill(start_color="C5A46D", end_color="C5A46D", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    # Encabezado
    ws.merge_cells("A1:J1")
    ws["A1"] = "CYREX STORE - REPORTE DE VENTAS"
    ws["A1"].font = Font(bold=True, size=14, color="171B24")
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:J2")
    rango = f"Período: {fecha_inicio or 'Inicio'} al {fecha_fin or 'Fin'}"
    ws["A2"] = rango
    ws["A2"].font = Font(size=10, italic=True)
    ws["A2"].alignment = Alignment(horizontal="center")

    # Fila 4: headers
    headers = ["N° Factura", "Fecha", "Cliente", "Correo", "Productos", "Subtotal", "Impuestos", "Descuentos", "Total Neto", "Método Pago"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    # Datos
    for i, row_data in enumerate(datos, 5):
        values = [
            row_data["numero_factura"], row_data["fecha"], row_data["cliente"],
            row_data["correo"], row_data["productos"], row_data["subtotal"],
            row_data["impuestos"], row_data["descuentos"], row_data["total_neto"],
            row_data["metodo_pago"],
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=i, column=col, value=val)
            cell.border = thin_border
            if col >= 6 and col <= 9:
                cell.number_format = '#,##0.00'

    # Fila resumen
    if datos:
        resumen_row = len(datos) + 6
        ws.cell(row=resumen_row, column=5, value="TOTALES:").font = Font(bold=True)
        ws.cell(row=resumen_row, column=6, value=sum(d["subtotal"] for d in datos)).font = Font(bold=True)
        ws.cell(row=resumen_row, column=6).number_format = '#,##0.00'
        ws.cell(row=resumen_row, column=7, value=sum(d["impuestos"] for d in datos)).font = Font(bold=True)
        ws.cell(row=resumen_row, column=7).number_format = '#,##0.00'
        ws.cell(row=resumen_row, column=8, value=sum(d["descuentos"] for d in datos)).font = Font(bold=True)
        ws.cell(row=resumen_row, column=8).number_format = '#,##0.00'
        ws.cell(row=resumen_row, column=9, value=sum(d["total_neto"] for d in datos)).font = Font(bold=True)
        ws.cell(row=resumen_row, column=9).number_format = '#,##0.00'

    # Ajustar anchos de columna
    widths = [18, 12, 25, 30, 40, 15, 15, 15, 15, 15]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"reporte_ventas_cyrex.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get(
    "/reporte/pdf",
    summary="Exportar reporte de ventas en PDF",
    responses={200: {"description": "Archivo PDF"}},
)
def reporte_pdf(
    fecha_inicio: str | None = Query(default=None, description="Formato: YYYY-MM-DD"),
    fecha_fin: str | None = Query(default=None, description="Formato: YYYY-MM-DD"),
    _: dict = Depends(require_roles("Administrador")),
    db: Session = Depends(get_db),
):
    rows = generar_datos_excel(db, fecha_inicio, fecha_fin)
    buffer = sales_report_pdf(rows, fecha_inicio, fecha_fin)
    filename = f"reporte_ventas_cyrex_{fecha_inicio or 'completo'}.pdf"
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get(
    "/factura/{numero_factura}/pdf",
    summary="Descargar factura en PDF",
    responses={200: {"description": "Factura PDF"}, 403: {"description": "Factura de otro cliente"}},
)
def factura_pdf(
    numero_factura: str,
    user: dict = Depends(current_user),
    db: Session = Depends(get_db),
):
    venta = db.scalar(select(Venta).where(Venta.numero_factura == numero_factura))
    if not venta:
        raise RecursoNoEncontrado("Factura", numero_factura)
    orden = db.get(Orden, venta.orden_id)
    is_staff = user["rol_nombre"] in {"Administrador", "Empleado"}
    if not orden or (not is_staff and orden.usuario_id != user["id"]):
        from fastapi import HTTPException
        raise HTTPException(status_code=403, detail="No puedes descargar esta factura")
    invoice = obtener_venta_por_numero(db, numero_factura)
    buffer = invoice_pdf(invoice)
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=factura_{numero_factura}.pdf"},
    )


@router.get(
    "/factura/{numero_factura}",
    summary="Buscar venta por número de factura",
    responses={200: {"description": "Venta encontrada"}, 404: {"description": "Factura no encontrada"}},
)
def get_by_numero(
    numero_factura: str,
    _: dict = Depends(require_roles("Administrador", "Empleado")),
    db: Session = Depends(get_db),
):
    venta = obtener_venta_por_numero(db, numero_factura)
    if not venta:
        raise RecursoNoEncontrado("Factura", numero_factura)
    return venta


@router.get(
    "/{venta_id}",
    summary="Obtener venta por ID",
    responses={200: {"description": "Venta encontrada"}, 404: {"description": "Venta no encontrada"}},
)
def get_venta_by_id(
    venta_id: int,
    _: dict = Depends(require_roles("Administrador", "Empleado")),
    db: Session = Depends(get_db),
):
    venta = obtener_venta_por_id(db, venta_id)
    if not venta:
        raise RecursoNoEncontrado("Venta", venta_id)
    return venta


@router.patch(
    "/{venta_id}",
    summary="Actualizar venta",
    responses={200: {"description": "Venta actualizada"}, 404: {"description": "Venta no encontrada"}},
)
def update_venta(
    venta_id: int,
    data: VentaActualizar,
    _: dict = Depends(require_roles("Administrador")),
    db: Session = Depends(get_db),
):
    venta = actualizar_venta(
        db,
        venta_id,
        notas_factura=data.notas_factura,
        descuento_porcentaje=data.descuento_porcentaje,
        estado=data.estado,
    )
    if not venta:
        raise RecursoNoEncontrado("Venta", venta_id)
    return {"message": "Venta actualizada", "venta": venta}
